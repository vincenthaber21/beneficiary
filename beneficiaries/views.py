import io
import json
import re
from types import SimpleNamespace
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.db.models import Q, Count
from django.db import IntegrityError, transaction
from django.http import HttpResponse, JsonResponse
from django.urls import reverse
from django.utils import timezone
from django.utils.safestring import mark_safe
from datetime import timedelta
from utils import rfid_access_required, staff_required
from checker.logic import evaluate, VALID_CLASSES, INCOME_THRESHOLD
from records.models import GrantRecord
from .models import Beneficiary, FamilyMember
from .forms import (
    BeneficiaryForm,
    FamilyMemberFormSet,
    available_citizen_count,
    invalidate_citizen_enrollment_cache,
    make_beneficiary_fingerprint_formset,
    search_available_citizens,
    search_beneficiaries,
)
from .household import household_cooldown, family_block_for


def _get_family_groups(beneficiaries):
    """Return beneficiaries grouped into family units via FamilyMember links.
    Each group is a dict: {label, members, is_family}.
    Multi-member groups are sorted first."""
    pk_to_ben = {b.pk: b for b in beneficiaries}
    ben_pks = set(pk_to_ben)

    adj = {pk: set() for pk in ben_pks}
    for ben_id, mem_id in (
        FamilyMember.objects
        .filter(beneficiary_id__in=ben_pks, member_id__in=ben_pks)
        .exclude(member__isnull=True)
        .values_list("beneficiary_id", "member_id")
    ):
        adj[ben_id].add(mem_id)
        adj[mem_id].add(ben_id)

    visited = set()
    raw_groups = []
    for b in beneficiaries:          # preserves original ordering
        if b.pk in visited:
            continue
        component = []
        stack = [b.pk]
        while stack:
            cur = stack.pop()
            if cur in visited:
                continue
            visited.add(cur)
            component.append(pk_to_ben[cur])
            stack.extend(adj[cur] - visited)
        raw_groups.append(component)

    # Multi-member families first, then individuals; each tier sorted by last name
    raw_groups.sort(key=lambda g: (len(g) == 1, g[0].last_name, g[0].first_name))

    result = []
    for group in raw_groups:
        last_names = sorted({b.last_name for b in group})
        label = " & ".join(last_names[:2])
        if len(last_names) > 2:
            label += f" +{len(last_names) - 2}"
        result.append({"label": label, "members": group, "is_family": len(group) > 1})
    return result


def _attach_eligibility(queryset):
    """Attach an ``eligibility_result`` dict to each Beneficiary in the queryset
    using a single extra query to find recent grants (avoids N+1). Also enforces
    the household rule: a grant to any family member blocks the whole family."""
    cutoff = timezone.now().date() - timedelta(days=90)
    recent_ids = set(
        GrantRecord.objects.filter(
            status="released", date_granted__gte=cutoff
        ).values_list("beneficiary_id", flat=True)
    )
    blocker_id_map, name_map, _ = household_cooldown(3)
    result_list = []
    for b in queryset:
        # Check both system grant records AND manually entered prior_grant_date
        has_recent = (
            b.pk in recent_ids
            or (b.prior_grant_date is not None and b.prior_grant_date >= cutoff)
        )
        family_blocker = None if has_recent else family_block_for(b.pk, blocker_id_map, name_map)
        b.family_blocker = family_blocker
        b.eligibility_result = evaluate(
            beneficiary_class=b.beneficiary_class,
            received_grant_within_3_months=has_recent,
            monthly_income=float(b.monthly_income),
            household_blocker_name=family_blocker,
        )
        result_list.append(b)
    return result_list


def _filtered_beneficiary_queryset(request):
    """Apply list-page filters (q, class, family, claim) and return
    (queryset, q, cls, fam, claim)."""
    q = request.GET.get("q", "").strip()
    cls = request.GET.get("class", "")
    fam = request.GET.get("family", "")
    claim = request.GET.get("claim", "").strip()
    qs = Beneficiary.objects.select_related("citizen").all()
    if q:
        qs = qs.filter(
            Q(last_name__icontains=q) | Q(first_name__icontains=q) |
            Q(rfid_id__icontains=q) | Q(contact_number__icontains=q) |
            Q(citizen__registry_no__icontains=q)
        )
    if cls:
        qs = qs.filter(beneficiary_class=cls)

    claimed_ids = set(
        GrantRecord.objects.filter(status="released")
        .values_list("beneficiary_id", flat=True)
        .distinct()
    )
    if claim == "claimed":
        qs = qs.filter(pk__in=claimed_ids)
    elif claim == "not_claimed":
        qs = qs.exclude(pk__in=claimed_ids)

    # Beneficiaries that participate in at least one family link (either direction)
    linked_ids = set()
    for a, b in (
        FamilyMember.objects.exclude(member__isnull=True)
        .values_list("beneficiary_id", "member_id")
    ):
        linked_ids.add(a)
        linked_ids.add(b)
    if fam == "family":
        qs = qs.filter(pk__in=linked_ids)
    elif fam == "individual":
        qs = qs.exclude(pk__in=linked_ids)

    return qs, q, cls, fam, claim


def _annotate_claim_status(beneficiaries):
    """Attach received_count / has_received from released grant records."""
    received_map = dict(
        GrantRecord.objects.filter(status="released")
        .values_list("beneficiary_id")
        .annotate(n=Count("id"))
        .values_list("beneficiary_id", "n")
    )
    for b in beneficiaries:
        b.received_count = received_map.get(b.pk, 0)
        b.has_received = b.received_count > 0
    return beneficiaries


@staff_required
def beneficiary_list(request):
    qs, q, cls, fam, claim = _filtered_beneficiary_queryset(request)
    beneficiaries = _annotate_claim_status(_attach_eligibility(qs))

    family_groups = _get_family_groups(beneficiaries)
    family_count = sum(1 for g in family_groups if g["is_family"])
    received_count = sum(1 for b in beneficiaries if b.has_received)
    not_claimed_count = sum(1 for b in beneficiaries if not b.has_received)
    return render(request, "beneficiaries/list.html", {
        "beneficiaries": beneficiaries,
        "family_groups": family_groups,
        "total": len(beneficiaries),
        "family_count": family_count,
        "claimed_count": received_count,
        "not_claimed_count": not_claimed_count,
        "q": q, "cls": cls, "fam": fam, "claim": claim,
        "classes": VALID_CLASSES,
    })


@staff_required
def beneficiary_export_excel(request):
    """Download the currently filtered beneficiary list as an .xlsx file."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, "Excel export requires openpyxl. Run: pip install openpyxl")
        return redirect("beneficiaries:list")

    qs, q, cls, fam, claim = _filtered_beneficiary_queryset(request)
    beneficiaries = _annotate_claim_status(_attach_eligibility(qs))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Beneficiaries"

    headers = [
        "ID", "Last Name", "First Name", "Middle Name", "Sex",
        "Class", "Monthly Income", "Eligibility", "Eligibility Reason",
        "Claim Status", "Claims Count",
        "Contact", "Address", "Barangay", "Municipality", "Province",
        "RFID / ID", "Prior Grant Date", "Date Registered", "Status", "Notes",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="B80E1C")
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    sex_labels = dict(Beneficiary.SEX_CHOICES)
    for row_idx, b in enumerate(beneficiaries, 2):
        elig = getattr(b, "eligibility_result", {}) or {}
        claim_label = "Claimed" if b.has_received else "Does not claim"
        values = [
            b.pk,
            b.last_name,
            b.first_name,
            b.middle_name or "",
            sex_labels.get(b.sex, b.sex or ""),
            b.beneficiary_class,
            float(b.monthly_income),
            "Qualified" if elig.get("qualified") else "Not Qualified",
            elig.get("reason", ""),
            claim_label,
            b.received_count,
            b.contact_number or "",
            b.address or "",
            b.barangay or "",
            b.municipality or "",
            b.province or "",
            b.rfid_id or "",
            b.prior_grant_date.isoformat() if b.prior_grant_date else "",
            b.date_registered.isoformat() if b.date_registered else "",
            "Active" if b.is_active else "Inactive",
            b.notes or "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin
            if col == 7:
                cell.number_format = '#,##0.00'

    # Readable column widths
    widths = {
        1: 8, 2: 16, 3: 16, 4: 14, 5: 8, 6: 18, 7: 14, 8: 14, 9: 40,
        10: 14, 11: 12, 12: 14, 13: 28, 14: 14, 15: 14, 16: 14,
        17: 16, 18: 14, 19: 14, 20: 10, 21: 30,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    # Filename reflects active filters when possible
    parts = ["beneficiaries"]
    if claim == "not_claimed":
        parts.append("not_claimed")
    elif claim == "claimed":
        parts.append("claimed")
    if cls:
        parts.append(cls.replace(" ", "_")[:30])
    if fam:
        parts.append(fam)
    if q:
        parts.append("search")
    filename = "_".join(parts) + ".xlsx"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@staff_required
def beneficiary_create(request):
    initial = {}
    preselect = request.GET.get("citizen", "").strip()
    if preselect.isdigit():
        initial["citizen"] = int(preselect)

    form = BeneficiaryForm(request.POST or None, initial=initial)
    formset = FamilyMemberFormSet(request.POST or None)
    fingerprint_formset = make_beneficiary_fingerprint_formset(
        request.POST or None, prefix="fp"
    )
    if (
        request.method == "POST"
        and form.is_valid()
        and formset.is_valid()
        and fingerprint_formset.is_valid()
    ):
        try:
            with transaction.atomic():
                ben = form.save()
                formset.instance = ben
                formset.save()
                fingerprint_formset.instance = ben
                fingerprint_formset.save()
        except IntegrityError as exc:
            msg = str(exc).lower()
            if "citizen" in msg or "uniq" in msg or "rfid" in msg:
                messages.error(
                    request,
                    "Could not enroll this citizen — they may already be a beneficiary, "
                    "or their RFID / ID conflicts with another record.",
                )
            else:
                messages.error(
                    request,
                    "A fingerprint you scanned is already registered to another person. "
                    "Each finger must be unique.",
                )
        else:
            # Auto-validate: new beneficiary has no system grant records yet;
            # received_grant_within_months already checks prior_grant_date field.
            personal_recent = ben.received_grant_within_months(3)
            blocker_id_map, name_map, _ = household_cooldown(3)
            family_blocker = None if personal_recent else family_block_for(ben.pk, blocker_id_map, name_map)
            result = evaluate(
                beneficiary_class=ben.beneficiary_class,
                received_grant_within_3_months=personal_recent,
                monthly_income=float(ben.monthly_income),
                household_blocker_name=family_blocker,
            )
            if result["qualified"]:
                messages.success(request, f"Beneficiary added. Eligibility: QUALIFIED — {result['reason']}")
            else:
                messages.warning(request, f"Beneficiary added. Eligibility: NOT QUALIFIED — {result['reason']}")
            invalidate_citizen_enrollment_cache()
            return redirect("beneficiaries:detail", pk=ben.pk)

    available_count = available_citizen_count()
    return render(request, "beneficiaries/form.html", {
        "form": form, "formset": formset, "action": "Add",
        "fingerprint_formset": fingerprint_formset,
        "card": _build_form_card(),
        "valid_classes_json": json.dumps(VALID_CLASSES),
        "income_threshold": INCOME_THRESHOLD,
        "citizen_lookup_url": reverse("beneficiaries:citizen_lookup"),
        "available_citizen_count": available_count,
        "enrollable_citizen_count": available_count,
        "citizen_create_url": reverse("citizens:create")
        + "?next="
        + reverse("beneficiaries:create"),
        "beneficiary_search_url": reverse("beneficiaries:search"),
        "exclude_beneficiary_pk": "",
    })


@staff_required
def beneficiary_update(request, pk):
    ben = get_object_or_404(Beneficiary.objects.select_related("citizen"), pk=pk)
    if request.method != "POST" and ben.citizen_id:
        # Keep personal cache in sync with citizen registry before editing.
        ben.sync_from_citizen()
    form = BeneficiaryForm(request.POST or None, instance=ben)
    formset = FamilyMemberFormSet(request.POST or None, instance=ben)
    fingerprint_formset = make_beneficiary_fingerprint_formset(
        request.POST or None, instance=ben, prefix="fp"
    )
    if (
        request.method == "POST"
        and form.is_valid()
        and formset.is_valid()
        and fingerprint_formset.is_valid()
    ):
        try:
            with transaction.atomic():
                ben = form.save()
                formset.save()
                fingerprint_formset.save()
        except IntegrityError:
            messages.error(
                request,
                "A fingerprint you scanned is already registered to another person. "
                "Each finger must be unique.",
            )
        else:
            personal_recent = ben.received_grant_within_months(3)
            blocker_id_map, name_map, _ = household_cooldown(3)
            family_blocker = None if personal_recent else family_block_for(ben.pk, blocker_id_map, name_map)
            result = evaluate(
                beneficiary_class=ben.beneficiary_class,
                received_grant_within_3_months=personal_recent,
                monthly_income=float(ben.monthly_income),
                household_blocker_name=family_blocker,
            )
            if result["qualified"]:
                messages.success(request, f"Beneficiary updated. Eligibility: QUALIFIED — {result['reason']}")
            else:
                messages.warning(request, f"Beneficiary updated. Eligibility: NOT QUALIFIED — {result['reason']}")
            return redirect("beneficiaries:detail", pk=pk)
    # Most recent RELEASED grant on file (system record) — the edit page's live
    # preview must account for this, not just the manual prior_grant_date field.
    last_system_grant = (
        ben.grants.filter(status="released")
        .order_by("-date_granted")
        .values_list("date_granted", flat=True)
        .first()
    )

    # Household rule – a grant to any family member blocks the whole family.
    # Only relevant when this person hasn't personally received one recently.
    personal_recent = ben.received_grant_within_months(3)
    blocker_id_map, name_map, _ = household_cooldown(3)
    family_blocker = None if personal_recent else family_block_for(ben.pk, blocker_id_map, name_map)

    # Family connections in BOTH directions:
    #  • forward  → people listed as this person's family members
    #  • reverse  → other beneficiaries who list this person as their family member
    family_connections = []
    seen_person_ids = set()
    for fm in ben.family_members.filter(member__isnull=False).select_related("member"):
        family_connections.append({"person": fm.member, "relationship": fm.relationship, "age": fm.age})
        seen_person_ids.add(fm.member_id)
    for fm in ben.member_of.filter(member__isnull=False).select_related("beneficiary"):
        if fm.beneficiary_id in seen_person_ids:
            continue
        family_connections.append({"person": fm.beneficiary, "relationship": fm.relationship, "age": fm.age})
        seen_person_ids.add(fm.beneficiary_id)
    return render(request, "beneficiaries/form.html", {
        "form": form, "formset": formset, "action": "Edit", "beneficiary": ben,
        "fingerprint_formset": fingerprint_formset,
        "card": _build_form_card(ben),
        "existing_family_members": family_connections,
        "last_system_grant": last_system_grant.isoformat() if last_system_grant else "",
        "family_blocker": family_blocker or "",
        "valid_classes_json": json.dumps(VALID_CLASSES),
        "income_threshold": INCOME_THRESHOLD,
        "citizen_lookup_url": reverse("beneficiaries:citizen_lookup"),
        "available_citizen_count": available_citizen_count(ben.pk),
        "enrollable_citizen_count": available_citizen_count(ben.pk),
        "beneficiary_search_url": reverse("beneficiaries:search"),
        "exclude_beneficiary_pk": ben.pk,
    })


@staff_required
def beneficiary_search(request):
    """JSON typeahead for linking family members on the enroll form."""
    q = request.GET.get("q", "").strip()
    exclude = request.GET.get("exclude", "").strip()
    exclude_pk = int(exclude) if exclude.isdigit() else None
    results = search_beneficiaries(q, exclude_pk=exclude_pk)
    return JsonResponse({
        "ok": True,
        "query": q,
        "count": len(results),
        "results": results,
    })


@staff_required
def citizen_lookup(request):
    """JSON: search citizens (?q=) or load one citizen (?id=) for the enroll form."""
    from django.http import JsonResponse
    from citizens.models import Citizen
    from .models import eligibility_from_citizen

    exclude = request.GET.get("exclude_beneficiary", "").strip()
    exclude_pk = int(exclude) if exclude.isdigit() else None

    # ── Search mode (AJAX typeahead) ─────────────────────────────────────
    q = request.GET.get("q", "").strip()
    if q:
        results = search_available_citizens(q, exclude_beneficiary_pk=exclude_pk)
        return JsonResponse({
            "ok": True,
            "query": q,
            "count": len(results),
            "results": results,
        })

    # ── Detail mode (fill personal fields + household) ───────────────────
    pk = request.GET.get("id", "").strip()
    if not pk.isdigit():
        return JsonResponse(
            {"ok": False, "error": "Provide ?q= to search or ?id= to load a citizen."},
            status=400,
        )
    try:
        c = Citizen.objects.get(pk=int(pk))
    except Citizen.DoesNotExist:
        return JsonResponse({"ok": False, "error": "Citizen not found."}, status=404)

    already = Beneficiary.objects.filter(citizen=c).select_related("citizen").first()
    if exclude_pk and already and already.pk == exclude_pk:
        already = None

    suggested_class, suggested_income = eligibility_from_citizen(c)
    income_range_display = ""
    if c.income_range:
        income_range_display = dict(c.INCOME_RANGE_CHOICES).get(c.income_range, c.income_range)

    payload = {
        "ok": True,
        "already_enrolled": bool(already),
        "citizen": {
            "id": c.pk,
            "registry_no": c.registry_no or "",
            "last_name": c.last_name,
            "first_name": c.first_name,
            "middle_name": c.middle_name or "",
            "sex": c.sex or "",
            "sex_display": c.get_sex_display() if c.sex else "",
            "contact_number": c.contact_number or "",
            "address": c.address or "",
            "barangay": c.barangay or "",
            "municipality": c.municipality or "",
            "province": c.province or "",
            "rfid_id": c.rfid_tag or c.registry_no or "",
            "signature": c.signature or "",
            "fingerprint_count": c.fingerprints.count(),
            "income_range": c.income_range or "",
            "income_range_display": income_range_display,
            "suggested_monthly_income": suggested_income,
            "suggested_class": suggested_class,
            "edit_url": f"/citizens/{c.pk}/edit/",
            "family_members": _citizen_household_for_enroll(c, exclude_beneficiary_pk=exclude_pk),
        },
    }
    if already:
        payload["error"] = (
            f"{c.full_name} is already enrolled as a beneficiary. "
            "Open their beneficiary record instead of enrolling again."
        )
        payload["beneficiary"] = {
            "id": already.pk,
            "detail_url": f"/beneficiaries/{already.pk}/",
            "class": already.beneficiary_class,
        }
    return JsonResponse(payload)


_BEN_RELATIONSHIPS = {
    "Spouse", "Child", "Parent", "Sibling", "Grandparent", "Grandchild", "Other",
}
_REL_INVERT = {
    "Spouse": "Spouse",
    "Child": "Parent",
    "Parent": "Child",
    "Sibling": "Sibling",
    "Grandparent": "Grandchild",
    "Grandchild": "Grandparent",
    "In-law": "Other",
    "Other": "Other",
}


def _map_ben_relationship(raw, invert=False):
    rel = (raw or "").strip()
    if invert:
        rel = _REL_INVERT.get(rel, "Other")
    if rel == "In-law":
        rel = "Other"
    if rel in _BEN_RELATIONSHIPS:
        return rel
    return "Other" if rel else ""


def _age_from_dob(dob, today=None):
    if not dob:
        return None
    today = today or timezone.now().date()
    years = today.year - dob.year
    if (today.month, today.day) < (dob.month, dob.day):
        years -= 1
    return max(0, years)


def _citizen_household_for_enroll(citizen, exclude_beneficiary_pk=None):
    """
    Build household rows from Citizen Family Composition for the enroll form.
    Enrolled relatives can be linked as beneficiary family members automatically.
    """
    from citizens.models import CitizenFamilyMember

    today = timezone.now().date()
    rows = []
    seen_keys = set()

    forward = list(
        citizen.family_members.select_related("member").all()
    )
    reverse = list(
        CitizenFamilyMember.objects
        .filter(member=citizen)
        .select_related("citizen")
    )

    related_citizen_ids = set()
    for fm in forward:
        if fm.member_id:
            related_citizen_ids.add(fm.member_id)
    for fm in reverse:
        if fm.citizen_id and fm.citizen_id != citizen.pk:
            related_citizen_ids.add(fm.citizen_id)

    ben_by_citizen = {
        b.citizen_id: b
        for b in Beneficiary.objects.filter(citizen_id__in=related_citizen_ids)
    }

    def add_row(*, name, relationship, age, citizen_obj=None, invert_rel=False):
        ben = None
        if citizen_obj is not None:
            ben = ben_by_citizen.get(citizen_obj.pk)
            if exclude_beneficiary_pk and ben and ben.pk == exclude_beneficiary_pk:
                return
        key = (
            f"b:{ben.pk}" if ben else
            f"c:{citizen_obj.pk}" if citizen_obj else
            f"n:{(name or '').casefold()}"
        )
        if not name or key in seen_keys:
            return
        seen_keys.add(key)
        rel = _map_ben_relationship(relationship, invert=invert_rel)
        if age is None and citizen_obj is not None:
            age = _age_from_dob(citizen_obj.date_of_birth, today)
        rows.append({
            "name": name,
            "relationship": rel or "Other",
            "age": age,
            "citizen_id": citizen_obj.pk if citizen_obj else None,
            "registry_no": (citizen_obj.registry_no if citizen_obj else "") or "",
            "beneficiary_id": ben.pk if ben else None,
            "beneficiary_label": (
                f"{ben.last_name}, {ben.first_name}"
                + (f" ({ben.beneficiary_class})" if ben and ben.beneficiary_class else "")
            ) if ben else "",
            "enrolled": bool(ben),
            "can_link": bool(ben),
        })

    for fm in forward:
        member = fm.member
        name = ""
        if member:
            name = member.full_name
        elif fm.full_name:
            name = fm.full_name.strip()
        add_row(
            name=name,
            relationship=fm.relationship,
            age=None,
            citizen_obj=member,
            invert_rel=False,
        )

    for fm in reverse:
        other = fm.citizen
        if not other or other.pk == citizen.pk:
            continue
        add_row(
            name=other.full_name,
            relationship=fm.relationship,
            age=None,
            citizen_obj=other,
            invert_rel=True,
        )

    return rows


@staff_required
def beneficiary_delete(request, pk):
    ben = get_object_or_404(Beneficiary, pk=pk)
    if request.method == "POST":
        ben.delete()
        invalidate_citizen_enrollment_cache()
        messages.success(request, "Beneficiary removed.")
        return redirect("beneficiaries:list")
    return render(request, "beneficiaries/confirm_delete.html", {"beneficiary": ben})


@staff_required
def beneficiary_detail(request, pk):
    ben = get_object_or_404(
        Beneficiary.objects.select_related("citizen"), pk=pk
    )
    grants = ben.grants.select_related("program", "granted_by").order_by("-date_granted")

    cutoff = timezone.now().date() - timedelta(days=90)

    # Step 1 – class check
    valid_class = ben.beneficiary_class in VALID_CLASSES

    # Step 2 – grant cooldown check (system records + prior_grant_date)
    last_system_grant = (
        ben.grants.filter(status="released")
        .order_by("-date_granted")
        .values_list("date_granted", flat=True)
        .first()
    )
    system_grant_in_cooldown = last_system_grant is not None and last_system_grant >= cutoff
    prior_grant_in_cooldown = (
        ben.prior_grant_date is not None and ben.prior_grant_date >= cutoff
    )
    received_recent = system_grant_in_cooldown or prior_grant_in_cooldown

    # Determine the effective "most recent grant date" that is in cooldown
    cooldown_source_date = None
    if system_grant_in_cooldown and prior_grant_in_cooldown:
        cooldown_source_date = max(last_system_grant, ben.prior_grant_date)
    elif system_grant_in_cooldown:
        cooldown_source_date = last_system_grant
    elif prior_grant_in_cooldown:
        cooldown_source_date = ben.prior_grant_date
    cooldown_end = cooldown_source_date + timedelta(days=90) if cooldown_source_date else None

    # Step 3 – income check
    income_ok = float(ben.monthly_income) < INCOME_THRESHOLD

    # Household rule – a grant to any family member blocks the whole family
    blocker_id_map, name_map, _ = household_cooldown(3)
    family_blocker = None if received_recent else family_block_for(ben.pk, blocker_id_map, name_map)

    result = evaluate(
        beneficiary_class=ben.beneficiary_class,
        received_grant_within_3_months=received_recent,
        monthly_income=float(ben.monthly_income),
        household_blocker_name=family_blocker,
    )
    # Family connections (both directions)
    family_members = (
        ben.family_members
        .filter(member__isnull=False)
        .select_related("member")
        .order_by("relationship", "member__last_name")
    )
    also_in_families = (
        ben.member_of
        .filter(member__isnull=False)
        .select_related("beneficiary")
        .order_by("relationship", "beneficiary__last_name")
    )
    return render(request, "beneficiaries/detail.html", {
        "beneficiary": ben,
        "grants": grants,
        "result": result,
        "family_members": family_members,
        "also_in_families": also_in_families,
        # Decision-tree step context
        "valid_class": valid_class,
        "received_recent": received_recent,
        "system_grant_in_cooldown": system_grant_in_cooldown,
        "prior_grant_in_cooldown": prior_grant_in_cooldown,
        "cooldown_source_date": cooldown_source_date,
        "cooldown_end": cooldown_end,
        "family_blocker": family_blocker,
        "income_ok": income_ok,
        "income_threshold": INCOME_THRESHOLD,
    })


def _qr_svg(data):
    """Return an inline-embeddable SVG string for a QR code of ``data``.
    Uses qrcode's pure-Python SVG factory so no image library is required.
    The width/height are stripped so the card CSS controls sizing."""
    import qrcode
    import qrcode.image.svg as svg
    qr = qrcode.QRCode(box_size=10, border=1, error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(image_factory=svg.SvgPathImage)
    buf = io.BytesIO()
    img.save(buf)
    out = buf.getvalue().decode()
    out = re.sub(r"<\?xml[^>]*\?>\s*", "", out)          # drop XML declaration
    out = re.sub(r'\swidth="[^"]*"', "", out, count=1)   # let CSS size it
    out = re.sub(r'\sheight="[^"]*"', "", out, count=1)
    return mark_safe(out)


@staff_required
def generate_id(request, pk):
    """Assign (or return the existing) e-Bahagi ID number for a beneficiary and
    return it as JSON, so the ID / Fingerprint Tap page can display it inline
    without navigating away. Format: EB-<year>-<zero-padded pk>."""
    from django.http import JsonResponse
    ben = get_object_or_404(Beneficiary, pk=pk)
    had_id = bool(ben.rfid_id)
    id_number = ben.ensure_id_number()
    return JsonResponse({"id_number": id_number, "newly_generated": not had_id})


def _build_card(ben):
    """Build the context dict for one printable e-Bahagi ID card. Auto-issues an
    ID number if one hasn't been assigned yet and embeds a scannable QR code."""
    id_number = ben.ensure_id_number()
    address_line = ", ".join(
        p for p in (ben.address, ben.barangay, ben.municipality) if p
    )
    return {
        "beneficiary": ben,
        "id_number": id_number,
        "address_line": address_line,
        "qr_svg": _qr_svg(id_number),
    }


def _build_form_card(ben=None):
    """Live preview card for the add/edit form — never auto-issues an ID."""
    if ben and ben.pk:
        address_line = ", ".join(
            p for p in (ben.address, ben.barangay, ben.municipality) if p
        )
        id_number = ben.rfid_id or "—"
        return {
            "beneficiary": ben,
            "id_number": id_number,
            "address_line": address_line,
            "qr_svg": _qr_svg(id_number) if ben.rfid_id else None,
        }
    return {
        "beneficiary": SimpleNamespace(
            last_name=getattr(ben, "last_name", "") if ben else "",
            first_name=getattr(ben, "first_name", "") if ben else "",
            middle_name=getattr(ben, "middle_name", "") if ben else "",
            date_of_birth=None,
            signature=getattr(ben, "signature", "") if ben else "",
        ),
        "id_number": "—",
        "address_line": "",
        "qr_svg": None,
    }


@staff_required
def beneficiary_id_card(request, pk):
    """Render a printable e-Bahagi ID card for a single beneficiary."""
    ben = get_object_or_404(Beneficiary, pk=pk)
    return render(request, "beneficiaries/id_card.html", {"card": _build_card(ben)})


@staff_required
def beneficiary_id_cards_all(request):
    """Render printable e-Bahagi ID cards for every beneficiary, one per page."""
    cards = [
        _build_card(ben)
        for ben in Beneficiary.objects.order_by("last_name", "first_name")
    ]
    return render(request, "beneficiaries/id_cards_all.html", {"cards": cards})


def _rfid_qualify(beneficiary):
    """Run eligibility check for a single beneficiary (shared by ID / name lookup)."""
    personal_recent = beneficiary.received_grant_within_months(3)
    blocker_id_map, name_map, _ = household_cooldown(3)
    family_blocker = (
        None if personal_recent
        else family_block_for(beneficiary.pk, blocker_id_map, name_map)
    )
    return evaluate(
        beneficiary_class=beneficiary.beneficiary_class,
        received_grant_within_3_months=personal_recent,
        monthly_income=float(beneficiary.monthly_income),
        household_blocker_name=family_blocker,
    )


@rfid_access_required
def rfid_check(request):
    from citizens.fingerprint_match import (
        find_best_beneficiary_fingerprint,
        find_best_fingerprint,
    )
    from citizens.models import CitizenFingerprint
    from .models import BeneficiaryFingerprint

    beneficiary = None
    result = None
    error = None
    rfid = ""
    name_matches = None
    first_name = ""
    middle_name = ""
    last_name = ""

    if request.method == "POST":
        # Pick a specific row from a prior name-search result list.
        pick_pk = request.POST.get("beneficiary_pk", "").strip()
        if pick_pk.isdigit():
            beneficiary = Beneficiary.objects.filter(pk=int(pick_pk)).first()
            if beneficiary:
                result = _rfid_qualify(beneficiary)
            else:
                error = "Selected beneficiary was not found."
        elif request.POST.get("search_by") == "name":
            first_name = request.POST.get("first_name", "").strip()
            middle_name = request.POST.get("middle_name", "").strip()
            last_name = request.POST.get("last_name", "").strip()
            if not (first_name or middle_name or last_name):
                error = "Enter at least a first, middle, or last name to search."
            else:
                qs = Beneficiary.objects.all()
                if first_name:
                    qs = qs.filter(first_name__icontains=first_name)
                if middle_name:
                    qs = qs.filter(middle_name__icontains=middle_name)
                if last_name:
                    qs = qs.filter(last_name__icontains=last_name)
                matches = list(qs.order_by("last_name", "first_name", "pk")[:25])
                if len(matches) == 1:
                    beneficiary = matches[0]
                    result = _rfid_qualify(beneficiary)
                elif len(matches) > 1:
                    name_matches = matches
                else:
                    parts = [p for p in (first_name, middle_name, last_name) if p]
                    error = f'No beneficiary found matching name "{" ".join(parts)}".'
        else:
            rfid = request.POST.get("rfid_id", "").strip()
            if rfid:
                try:
                    beneficiary = Beneficiary.objects.get(rfid_id=rfid)
                except Beneficiary.DoesNotExist:
                    # Exact / fuzzy match against enrolled beneficiary fingerprints.
                    bfp = (
                        BeneficiaryFingerprint.objects.select_related("beneficiary")
                        .filter(template_data=rfid)
                        .order_by("-is_primary", "-registered_at")
                        .first()
                    )
                    if not bfp:
                        bfp, _dist = find_best_beneficiary_fingerprint(rfid, threshold=10)
                    if bfp:
                        beneficiary = bfp.beneficiary
                    else:
                        # Fall back to citizen fingerprints, then map to a beneficiary.
                        fp = (
                            CitizenFingerprint.objects.select_related("citizen")
                            .filter(template_data=rfid)
                            .order_by("-is_primary", "-registered_at")
                            .first()
                        )
                        if not fp:
                            fp, _dist = find_best_fingerprint(rfid, threshold=10)

                        if fp:
                            citizen = fp.citizen
                            beneficiary = getattr(citizen, "beneficiary", None)
                            if beneficiary is None:
                                beneficiary = Beneficiary.objects.filter(citizen=citizen).first()
                            if not beneficiary:
                                # Legacy name / RFID fallback for unlinked records.
                                beneficiary = Beneficiary.objects.filter(
                                    last_name__iexact=citizen.last_name,
                                    first_name__iexact=citizen.first_name,
                                ).first()
                            if not beneficiary and citizen.rfid_tag:
                                beneficiary = Beneficiary.objects.filter(
                                    rfid_id=citizen.rfid_tag
                                ).first()
                            if not beneficiary:
                                error = (
                                    f'Fingerprint matched citizen "{citizen.full_name}" '
                                    f"({citizen.registry_no}), but no beneficiary record was found."
                                )
                        else:
                            error = f'No beneficiary found with ID / fingerprint "{rfid}".'

                if beneficiary:
                    result = _rfid_qualify(beneficiary)
            else:
                error = "Please enter or scan an ID / fingerprint, or search by name."

    return render(request, "beneficiaries/rfid.html", {
        "beneficiary": beneficiary,
        "result": result,
        "error": error,
        "rfid": rfid,
        "name_matches": name_matches,
        "first_name": first_name,
        "middle_name": middle_name,
        "last_name": last_name,
    })
