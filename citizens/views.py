import io
from types import SimpleNamespace

from django.contrib import messages
from django.core.exceptions import ObjectDoesNotExist
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from beneficiaries.models import class_from_citizen
from beneficiaries.views import _qr_svg
from checker.logic import VALID_CLASSES
from utils import staff_required

from .forms import CitizenForm, CitizenFamilyFormSet, make_citizen_fingerprint_formset
from .models import Citizen
from .search import filter_citizens_by_query, search_citizens

_CREATE_INITIAL = {
    "municipality": "Bacnotan",
    "province": "La Union",
}

# Map beneficiary-class labels → citizen vulnerable_groups / occupation keys
# (same derivation used by class_from_citizen).
_CLASS_FILTERS = {
    "Solo Parent": {"groups": ["solo_parent"]},
    "PWD": {"groups": ["pwd"]},
    "Senior Citizen": {"groups": ["senior"]},
    "Lactating Mother": {"groups": ["lactating"]},
    "Rice Farmer": {"groups": ["farmer"], "occupations": ["farmer"]},
    "Construction Worker": {"occupations": ["laborer"]},
    "Market Vendor": {"occupations": ["vendor"]},
    "Tricycle Driver": {"occupations": ["driver"]},
}


def _filter_by_vulnerable_class(qs, cls):
    """Filter citizens whose vulnerable group / occupation maps to ``cls``."""
    spec = _CLASS_FILTERS.get(cls)
    if not spec:
        return qs
    q_obj = Q()
    for g in spec.get("groups", []):
        # SQLite JSONField has no __contains; match the quoted key in stored JSON.
        q_obj |= Q(vulnerable_groups__icontains=f'"{g}"')
    for occ in spec.get("occupations", []):
        q_obj |= (
            Q(occupation_1=occ) | Q(occupation_2=occ) | Q(occupation_3=occ)
        )
    return qs.filter(q_obj) if q_obj else qs



def _build_card(citizen):
    """Context dict for the shared e-Bahagi ID card template.
    Auto-issues a Citizen ID (registry_no) when opening a printable card,
    matching beneficiary Generate ID behaviour."""
    if citizen and citizen.pk:
        id_number = citizen.ensure_id_number()
        address_line = ", ".join(
            p for p in (citizen.address, citizen.barangay, citizen.municipality) if p
        )
        return {
            "beneficiary": citizen,
            "id_number": id_number,
            "address_line": address_line,
            "qr_svg": _qr_svg(id_number),
        }
    return {
        "beneficiary": SimpleNamespace(
            last_name="",
            first_name="",
            middle_name="",
            date_of_birth=None,
            signature="",
        ),
        "id_number": "—",
        "address_line": "",
        "qr_svg": None,
    }


def _build_form_card(citizen=None):
    """Live preview card for the add/edit form — never auto-issues an ID."""
    if citizen and citizen.pk:
        address_line = ", ".join(
            p for p in (citizen.address, citizen.barangay, citizen.municipality) if p
        )
        id_number = citizen.registry_no or "—"
        return {
            "beneficiary": citizen,
            "id_number": id_number,
            "address_line": address_line,
            "qr_svg": _qr_svg(id_number) if citizen.registry_no else None,
        }
    return _build_card(None)


def _save_citizen_with_related(form, family_formset, fingerprint_formset):
    with transaction.atomic():
        citizen = form.save()
        family_formset.instance = citizen
        family_formset.save()
        fingerprint_formset.instance = citizen
        fingerprint_formset.save()
    return citizen


def _filtered_citizen_queryset(request):
    """Apply list-page filters and return (queryset, filter dict)."""
    q = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    barangay = request.GET.get("barangay", "").strip()
    aid = request.GET.get("aid", "").strip()  # enrolled | not_enrolled | availed | ""
    sex = request.GET.get("sex", "").strip().upper()
    cls = request.GET.get("class", "").strip()
    if cls not in VALID_CLASSES:
        cls = ""
    if sex not in ("M", "F"):
        sex = ""

    qs = Citizen.objects.select_related("beneficiary").all()
    if q:
        qs = filter_citizens_by_query(qs, q)
    if status:
        qs = qs.filter(status=status)
    if barangay:
        qs = qs.filter(barangay__iexact=barangay)
    if sex:
        qs = qs.filter(sex=sex)
    if cls:
        qs = _filter_by_vulnerable_class(qs, cls)
    if aid == "enrolled":
        qs = qs.filter(beneficiary__isnull=False)
    elif aid == "not_enrolled":
        qs = qs.filter(beneficiary__isnull=True)
    elif aid == "availed":
        # Enrolled and has at least one released grant, or a prior_grant_date on file.
        qs = qs.filter(beneficiary__isnull=False).filter(
            Q(beneficiary__grants__status="released")
            | Q(beneficiary__prior_grant_date__isnull=False)
        ).distinct()

    filters = {
        "q": q,
        "status": status,
        "barangay": barangay,
        "aid": aid,
        "sex": sex,
        "cls": cls,
    }
    return qs, filters


@staff_required
def citizen_list(request):
    qs, filters = _filtered_citizen_queryset(request)

    barangays = (
        Citizen.objects.exclude(barangay="")
        .values_list("barangay", flat=True)
        .distinct()
        .order_by("barangay")
    )

    enrolled_count = Citizen.objects.filter(beneficiary__isnull=False).count()
    total = qs.count()
    paginator = Paginator(qs, 50)
    page_obj = paginator.get_page(request.GET.get("page"))

    citizens = list(page_obj.object_list)
    for c in citizens:
        c.derived_class = class_from_citizen(c)

    context = {
        "citizens": citizens,
        "page_obj": page_obj,
        "total": total,
        "classes": VALID_CLASSES,
        "barangays": barangays,
        "status_choices": Citizen.STATUS_CHOICES,
        "count_active": Citizen.objects.filter(status="active").count(),
        "count_inactive": Citizen.objects.exclude(status="active").count(),
        "count_enrolled": enrolled_count,
        **filters,
    }
    return render(request, "citizens/list.html", context)


@staff_required
def citizen_export_excel(request):
    """Download the currently filtered citizen list as an .xlsx file."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, "Excel export requires openpyxl. Run: pip install openpyxl")
        return redirect("citizens:list")

    qs, filters = _filtered_citizen_queryset(request)
    citizens = qs.order_by("last_name", "first_name")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Citizens"

    headers = [
        "Citizen ID", "Last Name", "First Name", "Middle Name", "Suffix",
        "Sex", "Birth Date", "Place of Birth", "Civil Status",
        "Class", "Vulnerable Groups", "Occupation 1", "Occupation 2", "Occupation 3",
        "Contact", "Email", "Address", "Barangay", "Municipality", "Province", "Zip Code",
        "PhilSys ID", "Voter ID", "Other ID Type", "Other ID Number", "RFID Tag",
        "Education", "Income Range", "Health Conditions", "Religion",
        "Status", "Aid / Beneficiary", "Beneficiary Class", "Date Registered", "Notes",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="B80E1C")
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    occ_labels = dict(Citizen.OCCUPATION_CHOICES)
    edu_labels = dict(Citizen.EDUCATION_CHOICES)
    income_labels = dict(Citizen.INCOME_RANGE_CHOICES)
    civil_labels = dict(Citizen.CIVIL_STATUS_CHOICES)
    status_labels = dict(Citizen.STATUS_CHOICES)
    sex_labels = dict(Citizen.SEX_CHOICES)

    for row_idx, c in enumerate(citizens.iterator(chunk_size=500), 2):
        try:
            beneficiary = c.beneficiary
        except ObjectDoesNotExist:
            beneficiary = None
        if beneficiary is not None:
            aid_label = "Enrolled"
            ben_class = beneficiary.beneficiary_class or ""
        else:
            aid_label = "Not enrolled"
            ben_class = ""
        values = [
            c.registry_no or "",
            c.last_name,
            c.first_name,
            c.middle_name or "",
            c.suffix or "",
            sex_labels.get(c.sex, c.sex or ""),
            c.date_of_birth.isoformat() if c.date_of_birth else "",
            c.place_of_birth or "",
            civil_labels.get(c.civil_status, c.civil_status or ""),
            class_from_citizen(c) or "",
            ", ".join(c.vulnerable_labels()),
            occ_labels.get(c.occupation_1, c.occupation_1 or ""),
            occ_labels.get(c.occupation_2, c.occupation_2 or ""),
            occ_labels.get(c.occupation_3, c.occupation_3 or ""),
            c.contact_number or "",
            c.email or "",
            c.address or "",
            c.barangay or "",
            c.municipality or "",
            c.province or "",
            c.zip_code or "",
            c.philsys_id or "",
            c.voter_id or "",
            c.other_id_type or "",
            c.other_id_number or "",
            c.rfid_tag or "",
            edu_labels.get(c.education, c.education or ""),
            income_labels.get(c.income_range, c.income_range or ""),
            ", ".join(c.health_labels()),
            c.religion or "",
            status_labels.get(c.status, c.status or ""),
            aid_label,
            ben_class,
            c.date_registered.isoformat() if c.date_registered else "",
            c.notes or "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin

    widths = {
        1: 14, 2: 16, 3: 16, 4: 14, 5: 8, 6: 8, 7: 12, 8: 18, 9: 14,
        10: 16, 11: 24, 12: 16, 13: 16, 14: 16, 15: 14, 16: 22, 17: 24,
        18: 14, 19: 14, 20: 12, 21: 10, 22: 16, 23: 14, 24: 14, 25: 14,
        26: 14, 27: 18, 28: 18, 29: 22, 30: 14, 31: 12, 32: 14, 33: 16,
        34: 14, 35: 30,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    parts = ["citizens"]
    if filters["aid"]:
        parts.append(filters["aid"])
    if filters["cls"]:
        parts.append(filters["cls"].replace(" ", "_")[:30])
    if filters["status"]:
        parts.append(filters["status"])
    if filters["barangay"]:
        parts.append(filters["barangay"].replace(" ", "_")[:20])
    if filters["sex"]:
        parts.append(filters["sex"])
    if filters["q"]:
        parts.append("search")
    filename = "_".join(parts) + ".xlsx"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@staff_required
def citizen_detail(request, pk):
    citizen = get_object_or_404(
        Citizen.objects.select_related("beneficiary").prefetch_related(
            "family_members__member", "fingerprints", "beneficiary__grants"
        ),
        pk=pk,
    )
    beneficiary = None
    try:
        beneficiary = citizen.beneficiary
    except ObjectDoesNotExist:
        beneficiary = None
    last_grant = None
    grant_count = 0
    if beneficiary is not None:
        released = beneficiary.grants.filter(status="released").order_by("-date_granted")
        grant_count = released.count()
        last = released.first()
        last_grant = last.date_granted if last else beneficiary.prior_grant_date
    return render(
        request,
        "citizens/detail.html",
        {
            "citizen": citizen,
            "card": _build_card(citizen),
            "beneficiary": beneficiary,
            "last_grant": last_grant,
            "grant_count": grant_count,
            "has_availed": bool(last_grant) or grant_count > 0,
        },
    )


def _family_form_context(citizen=None):
    return {
        "citizen_search_url": reverse("citizens:search"),
        "exclude_citizen_pk": citizen.pk if citizen and citizen.pk else "",
    }


@staff_required
def citizen_search(request):
    """JSON typeahead for linking family members (cached)."""
    q = request.GET.get("q", "").strip()
    exclude = request.GET.get("exclude", "").strip()
    exclude_pk = int(exclude) if exclude.isdigit() else None
    results = search_citizens(q, exclude_pk=exclude_pk)
    return JsonResponse({
        "ok": True,
        "query": q,
        "count": len(results),
        "results": results,
    })


@staff_required
def citizen_create(request):
    form = CitizenForm(request.POST or None, request.FILES or None, initial=_CREATE_INITIAL)
    formset = CitizenFamilyFormSet(request.POST or None, prefix="family")
    fingerprint_formset = make_citizen_fingerprint_formset(
        request.POST or None, prefix="fp"
    )
    next_url = (request.GET.get("next") or request.POST.get("next") or "").strip()
    if (
        request.method == "POST"
        and form.is_valid()
        and formset.is_valid()
        and fingerprint_formset.is_valid()
    ):
        try:
            citizen = _save_citizen_with_related(form, formset, fingerprint_formset)
        except IntegrityError:
            messages.error(
                request,
                "A fingerprint you scanned is already registered to another person. "
                "Each finger must be unique.",
            )
        else:
            messages.success(
                request,
                f'Citizen "{citizen.full_name}" registered ({citizen.registry_no}).',
            )
            # Return to beneficiary enroll flow when requested.
            if next_url.startswith("/beneficiaries/add"):
                sep = "&" if "?" in next_url else "?"
                return redirect(f"{next_url}{sep}citizen={citizen.pk}")
            if next_url.startswith("/"):
                return redirect(next_url)
            return redirect("citizens:detail", pk=citizen.pk)
    return render(
        request,
        "citizens/form.html",
        {
            "form": form,
            "formset": formset,
            "fingerprint_formset": fingerprint_formset,
            "action": "Add",
            "citizen": None,
            "card": _build_form_card(None),
            "next_url": next_url,
            **_family_form_context(None),
        },
    )


@staff_required
def citizen_update(request, pk):
    citizen = get_object_or_404(Citizen, pk=pk)
    form = CitizenForm(
        request.POST or None, request.FILES or None, instance=citizen
    )
    formset = CitizenFamilyFormSet(
        request.POST or None, instance=citizen, prefix="family"
    )
    fingerprint_formset = make_citizen_fingerprint_formset(
        request.POST or None, instance=citizen, prefix="fp"
    )
    if (
        request.method == "POST"
        and form.is_valid()
        and formset.is_valid()
        and fingerprint_formset.is_valid()
    ):
        try:
            citizen = _save_citizen_with_related(form, formset, fingerprint_formset)
        except IntegrityError:
            messages.error(
                request,
                "A fingerprint you scanned is already registered to another person. "
                "Each finger must be unique.",
            )
        else:
            messages.success(request, f'Citizen "{citizen.full_name}" updated.')
            return redirect("citizens:detail", pk=citizen.pk)
    return render(
        request,
        "citizens/form.html",
        {
            "form": form,
            "formset": formset,
            "fingerprint_formset": fingerprint_formset,
            "action": "Edit",
            "citizen": citizen,
            "card": _build_form_card(citizen),
            **_family_form_context(citizen),
        },
    )


@staff_required
def generate_id(request, pk):
    """Assign (or return the existing) Citizen ID and return it as JSON."""
    citizen = get_object_or_404(Citizen, pk=pk)
    had_id = bool(citizen.registry_no)
    id_number = citizen.ensure_id_number()
    return JsonResponse({"id_number": id_number, "newly_generated": not had_id})


@staff_required
def citizen_id_card(request, pk):
    """Render a printable e-Bahagi ID card for a single citizen.
    Generates a Citizen ID if one has not been assigned yet."""
    citizen = get_object_or_404(Citizen, pk=pk)
    return render(
        request,
        "citizens/id_card.html",
        {"citizen": citizen, "card": _build_card(citizen)},
    )


@staff_required
def citizen_id_cards_all(request):
    """Render printable e-Bahagi ID cards for every citizen, one per page."""
    cards = [
        _build_card(citizen)
        for citizen in Citizen.objects.order_by("last_name", "first_name")
    ]
    return render(request, "citizens/id_cards_all.html", {"cards": cards})


@staff_required
def citizen_delete(request, pk):
    citizen = get_object_or_404(Citizen, pk=pk)
    if request.method == "POST":
        name = citizen.full_name
        citizen.delete()
        messages.success(request, f'Citizen "{name}" removed from the registry.')
        return redirect("citizens:list")
    return render(request, "citizens/confirm_delete.html", {"citizen": citizen})
